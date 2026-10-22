"""Build full, unfiltered prospect workbooks from the accepted store directories.

Master Prospects is every accepted-directory row. No page-count cutoff and no
Semrush traffic cutoff. Decision Makers are people already researched for those
stores. Excluded is the directory reject lists that carry a reason column.
"""
from __future__ import annotations

from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path(__file__).resolve().parent / "full"
PROSPECTS = Path("/tmp/prospects")
BRANCH = "cursor/gulf-prospect-research-473e"
RAW = f"https://github.com/abaszeinn-beep/research/raw/{BRANCH}/markets/prospects/full"

DARK = PatternFill("solid", fgColor="1F2933")
ORANGE = "F97316"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HEADER_BORDER = Border(bottom=Side(style="thick", color=ORANGE))
BODY_FONT = Font(name="Calibri", size=10)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
BAND = PatternFill("solid", fgColor="FFF7ED")
HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

# Reject sheets from the same directory workbooks as the accepted masters.
# Included only when the sheet has a reason column and at least one data row.
REJECT_SHEETS = (
    "Excluded",
    "Unverified Candidates",
    "Removed - platform subdomain",
    "StoreLeads gap - rejected",
)

EXPECTED = {
    "Saudi_Arabia_Ecommerce_FULL.xlsx": 3935,
    "Saudi_Arabia_Fashion_FULL.xlsx": 576,
    "UAE_Ecommerce_FULL.xlsx": 4772,
    "UAE_Fashion_FULL.xlsx": 371,
    "Qatar_Ecommerce_FULL.xlsx": 717,
    "Qatar_Fashion_FULL.xlsx": 124,
    "Kuwait_Ecommerce_FULL.xlsx": 502,
    "Kuwait_Fashion_FULL.xlsx": 124,
    "USA_Ecommerce_FULL.xlsx": 26213,
    "USA_Fashion_FULL.xlsx": 6726,
}

MARKETS = [
    {
        "file": "Saudi_Arabia_Ecommerce_FULL.xlsx",
        "master": PROSPECTS / "src_Saudi_Arabia_Ecommerce_Master_Directory.csv",
        "directory": Path("/tmp/m/w_saudi-ecommerce-directory-473e/Saudi_Arabia_Ecommerce_Master_Directory.xlsx"),
        "people": [PROSPECTS / "final_saudi/Saudi_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Saudi_Arabia_Ecommerce_Prospects.xlsx",
    },
    {
        "file": "Saudi_Arabia_Fashion_FULL.xlsx",
        "master": PROSPECTS / "src_Saudi_Arabia_Fashion_Stores.csv",
        "directory": Path("/tmp/m/w_fashion-stores-473e/markets/fashion/Saudi_Arabia_Fashion_Stores.xlsx"),
        "people": [PROSPECTS / "final_saudi/Saudi_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Saudi_Arabia_Fashion_Prospects.xlsx",
    },
    {
        "file": "UAE_Ecommerce_FULL.xlsx",
        "master": PROSPECTS / "src_UAE_Ecommerce_Master_Directory.csv",
        "directory": Path("/tmp/m/w_uae-ecommerce-directory-73c6/markets/uae/UAE_Ecommerce_Master_Directory.xlsx"),
        "people": [PROSPECTS / "final_gulf/UAE_people.csv"],
        "filtered": Path(__file__).resolve().parent / "UAE_Ecommerce_Prospects.xlsx",
    },
    {
        "file": "UAE_Fashion_FULL.xlsx",
        "master": PROSPECTS / "src_UAE_Fashion_Stores.csv",
        "directory": Path("/tmp/m/w_fashion-stores-473e/markets/fashion/UAE_Fashion_Stores.xlsx"),
        "people": [PROSPECTS / "final_gulf/UAE_people.csv"],
        "filtered": Path(__file__).resolve().parent / "UAE_Fashion_Prospects.xlsx",
    },
    {
        "file": "Qatar_Ecommerce_FULL.xlsx",
        "master": PROSPECTS / "src_Qatar_Ecommerce_Master_Directory.csv",
        "directory": Path("/tmp/m/w_qatar-ecommerce-directory-9c58/markets/qatar/Qatar_Ecommerce_Master_Directory.xlsx"),
        "people": [PROSPECTS / "final_gulf/Qatar_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Qatar_Ecommerce_Prospects.xlsx",
    },
    {
        "file": "Qatar_Fashion_FULL.xlsx",
        "master": PROSPECTS / "src_Qatar_Fashion_Stores.csv",
        "directory": Path("/tmp/m/w_fashion-stores-473e/markets/fashion/Qatar_Fashion_Stores.xlsx"),
        "people": [PROSPECTS / "final_gulf/Qatar_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Qatar_Fashion_Prospects.xlsx",
    },
    {
        "file": "Kuwait_Ecommerce_FULL.xlsx",
        "master": PROSPECTS / "src_Kuwait_Ecommerce_Master_Directory.csv",
        "directory": Path("/tmp/m/w_kuwait-ecommerce-directory-f1cc/markets/kuwait/Kuwait_Ecommerce_Master_Directory.xlsx"),
        "people": [PROSPECTS / "final_gulf/Kuwait_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Kuwait_Ecommerce_Prospects.xlsx",
    },
    {
        "file": "Kuwait_Fashion_FULL.xlsx",
        "master": PROSPECTS / "src_Kuwait_Fashion_Stores.csv",
        "directory": Path("/tmp/m/w_fashion-stores-473e/markets/fashion/Kuwait_Fashion_Stores.xlsx"),
        "people": [PROSPECTS / "final_gulf/Kuwait_people.csv"],
        "filtered": Path(__file__).resolve().parent / "Kuwait_Fashion_Prospects.xlsx",
    },
    {
        "file": "USA_Ecommerce_FULL.xlsx",
        "master": PROSPECTS / "src_USA_Ecommerce_Master_Directory.csv",
        "directory": Path("/tmp/m/w_usa-ecommerce-directory-fe40/markets/usa/USA_Ecommerce_Master_Directory.xlsx"),
        "people": [PROSPECTS / f"final_usa/USA_shard{i}_people.csv" for i in range(1, 7)],
        "filtered": Path(__file__).resolve().parent / "USA_Ecommerce_Prospects.xlsx",
    },
    {
        "file": "USA_Fashion_FULL.xlsx",
        "master": PROSPECTS / "src_USA_Fashion_Stores.csv",
        "directory": Path("/tmp/m/w_fashion-stores-473e/markets/fashion/USA_Fashion_Stores.xlsx"),
        "people": [PROSPECTS / f"final_usa/USA_shard{i}_people.csv" for i in range(1, 7)],
        "filtered": Path(__file__).resolve().parent / "USA_Fashion_Prospects.xlsx",
    },
]


def url_key(value) -> str:
    """Identity for joining a researched URL onto a directory URL.

    http/https, trailing slash, and a leading www are ignored. The path is kept,
    so distinct locale URLs stay distinct.
    """
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    if not text.lower().startswith(("http://", "https://")):
        text = "https://" + text
    parsed = urlparse(text)
    host = parsed.netloc.lower()
    if host.startswith("www."):
        host = host[4:]
    path = parsed.path.rstrip("/")
    return host + path


def page_is_blank(value) -> bool:
    text = "" if value is None else str(value).strip()
    return text == "" or text.lower() == "not found"


def as_number(value):
    text = "" if value is None else str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if number.is_integer():
        return int(number)
    return number


def passes_research_filter(count, traffic) -> str:
    """(page count blank or >= 500) AND (Semrush organic traffic >= 1000)."""
    if page_is_blank(count):
        count_ok = True
    else:
        count_num = as_number(count)
        count_ok = isinstance(count_num, (int, float)) and count_num >= 500
    traffic_num = as_number(traffic)
    traffic_ok = isinstance(traffic_num, (int, float)) and traffic_num >= 1000
    return "Yes" if count_ok and traffic_ok else "No"


def metric_column(name: str) -> bool:
    lowered = name.lower()
    return lowered.startswith("organic ") or "pages count" in lowered


def display_value(column: str, value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return None
    if metric_column(column):
        number = as_number(text)
        if number is not None and str(text).strip().replace(",", "").replace(".", "", 1).isdigit():
            return number
    return text


def is_single_url(value) -> bool:
    return isinstance(value, str) and value.startswith(("http://", "https://")) and " " not in value and ";" not in value


def find_columns(columns, *needles: str) -> str:
    for needle in needles:
        for column in columns:
            if needle in column.lower():
                return column
    raise KeyError(needles)


def read_master(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    before = len(frame)
    url_col = frame.columns[1]
    frame = frame.drop_duplicates(keep="first")
    dropped = before - len(frame)
    if dropped:
        print(f"  exact duplicate rows dropped from {path.name}: {dropped}")
    if frame[url_col].duplicated().any():
        dup = int(frame[url_col].duplicated().sum())
        print(f"  same URL kept where other columns differ: {dup}")
    return frame


def read_people(paths: list[Path]) -> pd.DataFrame:
    frames = [pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig") for path in paths]
    people = pd.concat(frames, ignore_index=True)
    before = len(people)
    people = people.drop_duplicates(keep="first")
    if len(people) != before:
        print(f"  exact duplicate people rows dropped: {before - len(people)}")
    return people


def read_reject_sheet(path: Path, sheet: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return None
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        wb.close()
        return None
    header = ["" if cell is None else str(cell).strip() for cell in header_row]
    while header and header[-1] == "":
        header.pop()
    if not header or not any("reason" in name.lower() for name in header):
        wb.close()
        return None
    data = []
    width = len(header)
    for row in rows:
        values = list(row[:width])
        if any(cell is not None and str(cell).strip() for cell in values):
            data.append(values)
    wb.close()
    if not data:
        return None
    return header, data


def load_excluded(path: Path):
    blocks = []
    for sheet in REJECT_SHEETS:
        found = read_reject_sheet(path, sheet)
        if found:
            blocks.append((sheet, found[0], found[1]))
    if not blocks:
        return [], []
    columns = ["Exclusion list"]
    seen = set(columns)
    for _, header, _ in blocks:
        for name in header:
            if name not in seen:
                columns.append(name)
                seen.add(name)
    rows = []
    for sheet, header, data in blocks:
        index = {name: pos for pos, name in enumerate(header)}
        for values in data:
            record = [None] * len(columns)
            record[0] = sheet
            for name, pos in index.items():
                cell = values[pos] if pos < len(values) else None
                if isinstance(cell, str):
                    cell = cell.strip() or None
                record[columns.index(name)] = cell
            rows.append(record)
    return columns, rows


def filtered_urls(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Master Prospects"]
    urls = []
    for index, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True)):
        if row[1]:
            urls.append(str(row[1]))
    wb.close()
    return urls


def style_sheet(ws: Worksheet, widths: list[int], numeric_cols: set[int]):
    for index, width in enumerate(widths, 1):
        cell = ws.cell(row=1, column=index)
        cell.fill = DARK
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(widths))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    ws.sheet_properties.tabColor = ORANGE
    if ws.max_row >= 2:
        ws.conditional_formatting.add(
            f"A2:{last_col}{ws.max_row}",
            FormulaRule(formula=["MOD(ROW(),2)=1"], fill=BAND),
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(widths)):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            if cell.column in numeric_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            elif is_single_url(cell.value):
                cell.hyperlink = cell.value
                cell.font = LINK_FONT


def column_width(name: str) -> int:
    lowered = name.lower()
    if "url" in lowered or "website" in lowered or "link" in lowered or "source" in lowered:
        return 42
    if "reason" in lowered or "position" in lowered or "method" in lowered or "note" in lowered:
        return 36
    return min(max(len(name) + 3, 16), 34)


def write_rows(ws: Worksheet, columns: list[str], rows: list[list]):
    ws.append(columns)
    numeric = set()
    for row in rows:
        ws.append(row)
    for index, column in enumerate(columns, 1):
        if metric_column(column) or "score" in column.lower() or column.lower().endswith("rank"):
            numeric.add(index)
    style_sheet(ws, [column_width(name) for name in columns], numeric)


def build_one(spec: dict) -> dict:
    print(f"\n{spec['file']}")
    print(f"  master source: {spec['master']}")
    master = read_master(spec["master"])
    columns = list(master.columns)
    count_col = find_columns(columns, "pages count", "website count")
    traffic_col = find_columns(columns, "organic traffic")
    url_col = columns[1]
    out_columns = columns + ["Passes Research Filter"]

    keys = []
    key_owner = {}
    collisions = 0
    rows = []
    yes = 0
    for record in master.itertuples(index=False, name=None):
        raw = dict(zip(columns, record))
        key = url_key(raw[url_col])
        if key in key_owner:
            collisions += 1
        else:
            key_owner[key] = len(rows)
        keys.append(key)
        flag = passes_research_filter(raw[count_col], raw[traffic_col])
        yes += flag == "Yes"
        rows.append([display_value(column, raw[column]) for column in columns] + [flag])
    if collisions:
        print(f"  URL-key collisions (distinct paths kept, www/slash only): {collisions}")

    people = read_people(spec["people"])
    people_cols = list(people.columns)
    link_col = "Website Link"
    matched = []
    unmatched = 0
    seen_people = set()
    link_index = people_cols.index(link_col)
    for record in people.itertuples(index=False, name=None):
        raw = dict(zip(people_cols, record))
        key = url_key(raw[link_col])
        if key not in key_owner:
            unmatched += 1
            continue
        values = [display_value(column, raw[column]) for column in people_cols]
        # Same person recorded against a slash or www variant of one directory URL.
        identity = (key,) + tuple(values[:link_index] + values[link_index + 1:])
        if identity in seen_people:
            continue
        seen_people.add(identity)
        matched.append((key_owner[key], values))
    matched.sort(key=lambda item: item[0])
    people_rows = [item[1] for item in matched]
    print(f"  people matched {len(people_rows)} / {len(people)} (unmatched stores not in this directory: {unmatched})")

    excluded_cols, excluded_rows = load_excluded(spec["directory"])
    print(f"  excluded rows {len(excluded_rows)} from {spec['directory'].name}")

    # Every previously published filtered store that is still in this directory
    # must be marked Yes. Stores absent from the directory are reported, not dropped.
    present = {key: rows[index][-1] for key, index in key_owner.items()}
    missing_from_full = []
    marked_no = []
    for url in filtered_urls(spec["filtered"]):
        key = url_key(url)
        if key not in present:
            missing_from_full.append(url)
        elif present[key] != "Yes":
            marked_no.append(url)
    if marked_no:
        print(f"  WARNING filtered workbook URLs marked No: {len(marked_no)} e.g. {marked_no[:5]}")
    if missing_from_full:
        print(f"  filtered workbook URLs not in this directory: {len(missing_from_full)}")

    expected = EXPECTED[spec["file"]]
    if len(rows) != expected:
        print(f"  COUNT DIFFERS from audit {expected}: actual {len(rows)} source {spec['master']}")

    wb = Workbook()
    master_ws = wb.active
    master_ws.title = "Master Prospects"
    write_rows(master_ws, out_columns, rows)
    people_ws = wb.create_sheet("Decision Makers")
    write_rows(people_ws, people_cols, people_rows)
    if excluded_rows:
        excluded_ws = wb.create_sheet("Excluded")
        write_rows(excluded_ws, excluded_cols, excluded_rows)

    OUT.mkdir(parents=True, exist_ok=True)
    destination = OUT / spec["file"]
    wb.save(destination)
    print(f"  wrote {destination.name}: master {len(rows)} yes {yes} people {len(people_rows)} excluded {len(excluded_rows)}")
    return {
        "file": spec["file"],
        "master": len(rows),
        "people": len(people_rows),
        "excluded": len(excluded_rows),
        "yes": yes,
        "expected": expected,
        "source": str(spec["master"]),
        "filtered_missing": len(missing_from_full),
        "filtered_marked_no": len(marked_no),
    }


def write_readme(results: list[dict]):
    lines = [
        "# Full prospect directories",
        "",
        "These workbooks contain every accepted store in the directory. They have no page-count cutoff and no Semrush traffic cutoff. `Passes Research Filter` is Yes only when the page count is blank (including `Not found`) or at least 500 and Semrush organic traffic is at least 1,000. Traffic that is missing or under 1,000 is No. Rows marked No stay in Master Prospects.",
        "",
        "Decision Makers are people already researched for that store. Stores with no research have no people row. Excluded collects the directory reject lists that include a reason (fashion exclusions, unverified candidates, removed platform subdomains, and StoreLeads rejections). Those rows are not in Master Prospects. The first Excluded column, Exclusion list, names the source list.",
        "",
        "| File | Master rows | Decision-maker rows | Excluded rows | Download |",
        "| --- | ---: | ---: | ---: | --- |",
    ]
    for result in results:
        url = f"{RAW}/{result['file']}"
        lines.append(
            f"| `{result['file']}` | {result['master']} | {result['people']} | {result['excluded']} | {url} |"
        )
    lines.append("")
    (OUT / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    results = [build_one(spec) for spec in MARKETS]
    write_readme(results)
    print("\nSUMMARY")
    for result in results:
        flag = "OK" if result["master"] == result["expected"] else "DIFF"
        print(
            f"{result['file']}\tmaster={result['master']}\tpeople={result['people']}\t"
            f"excluded={result['excluded']}\tyes={result['yes']}\t{flag}\t"
            f"filtered_missing={result['filtered_missing']}\tfiltered_no={result['filtered_marked_no']}"
        )


if __name__ == "__main__":
    main()

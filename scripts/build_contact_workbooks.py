#!/usr/bin/env python3
"""Build two-sheet partner workbooks from the partner lists and saved contacts."""
import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PARTNERS = "/workspace/data"
CONTACTS = "/opt/cursor/artifacts/research/contacts"
OUT_DIR = os.path.join(PARTNERS, "excel")

MARKETS = {
    "Lebanon": ["lebanon_01"],
    "Kuwait": ["kuwait_01"],
    "Qatar": ["qatar_01"],
    "UAE": ["uae_01", "uae_02"],
    "Saudi_Arabia": ["ksa_01", "ksa_02"],
    "United_States": [f"us_{i:02d}" for i in range(1, 21)],
}

COMPANY_HEADERS = [
    "Company Name",
    "Website",
    "Category",
    "Product Fit",
    "Emails",
    "Phone",
    "LinkedIn",
]
PEOPLE_HEADERS = [
    "Decision Maker Name",
    "Phone",
    "Email",
    "LinkedIn",
    "Position",
    "Company Name",
    "Website",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(vertical="top", wrap_text=True)


def load_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def website_key(value):
    text = (value or "").strip().lower()
    for prefix in ("https://", "http://"):
        if text.startswith(prefix):
            text = text[len(prefix):]
    if text.startswith("www."):
        text = text[4:]
    return text.rstrip("/")


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.row_dimensions[1].height = 22
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = "1F4E79"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"


def build(market, batches):
    partners = load_csv(os.path.join(PARTNERS, f"Autommerce_Partners_{market}.csv"))
    contacts = []
    people = []
    for batch in batches:
        contacts.extend(load_csv(os.path.join(CONTACTS, "company", f"{batch}.csv")))
        people.extend(load_csv(os.path.join(CONTACTS, "people", f"{batch}.csv")))

    by_site = {}
    for row in contacts:
        key = website_key(row.get("Website"))
        if key:
            by_site[key] = row

    company_rows = []
    enriched = 0
    for partner in partners:
        contact = by_site.get(website_key(partner.get("Website")), {})
        if contact.get("Emails") or contact.get("Phone") or contact.get("LinkedIn"):
            enriched += 1
        company_rows.append([
            partner.get("Company Name", ""),
            partner.get("Website", ""),
            partner.get("Category", ""),
            partner.get("Product Fit", ""),
            (contact.get("Emails") or "").strip(),
            (contact.get("Phone") or "").strip(),
            (contact.get("LinkedIn") or "").strip(),
        ])

    order = {website_key(row[1]): index for index, row in enumerate(company_rows)}
    seen = set()
    people_rows = []
    for person in people:
        name = (person.get("Decision Maker Name") or "").strip()
        if not name:
            continue
        site = website_key(person.get("Website"))
        identity = (name.lower(), site, (person.get("LinkedIn") or "").strip().lower())
        if identity in seen:
            continue
        seen.add(identity)
        people_rows.append([
            name,
            (person.get("Phone") or "").strip(),
            (person.get("Email") or "").strip(),
            (person.get("LinkedIn") or "").strip(),
            (person.get("Position") or "").strip(),
            (person.get("Company Name") or "").strip(),
            (person.get("Website") or "").strip(),
        ])
    people_rows.sort(key=lambda row: (order.get(website_key(row[6]), 10**6), row[0].lower()))

    wb = Workbook()
    companies = wb.active
    companies.title = "Company details"
    companies.append(COMPANY_HEADERS)
    for row in company_rows:
        companies.append(row)
    style_sheet(companies, [36, 42, 55, 16, 55, 42, 46])

    makers = wb.create_sheet("Decision makers")
    makers.append(PEOPLE_HEADERS)
    for row in people_rows:
        makers.append(row)
    style_sheet(makers, [32, 24, 36, 48, 42, 36, 42])

    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"Autommerce_Partners_{market}.xlsx")
    wb.save(path)
    return path, len(company_rows), enriched, len(people_rows)


def main():
    for market, batches in MARKETS.items():
        path, companies, enriched, people = build(market, batches)
        print(f"{os.path.basename(path)} companies={companies} with_contacts={enriched} people={people}")


if __name__ == "__main__":
    main()

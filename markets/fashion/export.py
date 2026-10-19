import json,collections,csv,random,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
S=json.load(open('stores.json'));E=json.load(open('excluded.json'))
qa=json.load(open('qa.json')) if len(sys.argv)>1 else {'n':0,'errors':0,'rate':None,'notes':''}
order={'Brand-Major':0,'Large':1,'Medium':2,'Small':3}
S.sort(key=lambda x:(order[x['tier']],x['name'].lower()))
H=['Store Name','Website','Type','Gender','Platform','Size Tier','Verification Status','Evidence URL']
def pl(p):
    p=(p or '').replace('custom/unknown','Custom/Unknown')
    return p.capitalize() if p.islower() and p!='Custom/Unknown' else p
rows=[[x['name'].replace('\n',' ').strip()[:80],x['site'],x['type'],x['gender'],pl(x['platform']) or 'Unknown',x['tier'],x['status'],x['ev']] for x in S]
wb=Workbook();ws=wb.active;ws.title='Fashion Stores';ws.append(H)
for r in rows: ws.append(r)
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78')
for col,w in zip('ABCDEFGH',[34,40,22,10,18,13,40,60]): ws.column_dimensions[col].width=w
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
w2=wb.create_sheet('Excluded');w2.append(['Store Name','Website','Reason','Directory Category','Source'])
for x in E: w2.append([x['name'].replace('\n',' ')[:80],x['site'],x['reason'],x['cat'],x['source']])
for c in w2[1]: c.font=Font(bold=True)
w3=wb.create_sheet('Coverage & Method')
def sec(t,cnt):
    w3.append([t]);w3.cell(w3.max_row,1).font=Font(bold=True)
    for k,v in sorted(cnt.items(),key=lambda x:-x[1]): w3.append(['  '+str(k),v])
    w3.append([])
w3.append(['UAE Fashion Stores – coverage & method']);w3.cell(1,1).font=Font(bold=True,size=14);w3.append(['Total fashion stores',len(S)]);w3.append(['Excluded rows',len(E)]);w3.append([])
sec('By channel (recall pass)',collections.Counter(x.get('channel','') for x in S));w3.append(['First-pass count (before recall pass)',289]);w3.append(['After recall pass',len(S)]);w3.append([]);sec('By source',collections.Counter(x['source'] for x in S));sec('By type',collections.Counter(x['type'] for x in S));sec('By gender',collections.Counter(x['gender'] for x in S))
sec('By platform',collections.Counter(pl(x['platform']) or 'Unknown' for x in S));sec('By size tier',collections.Counter(x['tier'] for x in S));sec('By verification status',collections.Counter(x['status'] for x in S))
sec('Exclusion reasons',collections.Counter(x['reason'] for x in E))
w3.append(['QA']);w3.cell(w3.max_row,1).font=Font(bold=True)
w3.append(['  Sample size',qa['n']]);w3.append(['  Errors (clear false positives)',qa['errors']]);w3.append(['  Error rate',qa['rate']]);w3.append(['  Notes',qa['notes']]);w3.append([])
for t in ['METHOD',
'1. Input: 4,690 unique verified Master Directory stores (incl. secondary-evidence and manually accepted majors) + 21,053 unverified candidates from branch cursor/uae-ecommerce-directory-73c6 (DB snapshot).',
'2. Live classification (aiohttp, <=2 conns/host, 15s timeout, no CAPTCHA/login bypass): homepage nav anchors + link paths + title/meta; Shopify /collections.json and /products.json product_type/tags; WooCommerce store API categories. English and Arabic clothing keywords (abaya, kandura, dresses, shirts, jeans, lingerie, activewear, modest, عبايات, فساتين, ملابس ...).',
'3. Rules (rules.py): directory category + live evidence; B2B uniform/workwear, tailoring-alteration services, kids-only, non-clothing categories with weak keyword hits are excluded. Gender from nav terms (women/men/abaya/kandura/lingerie). Type = Fashion-only if the directory category is fashion, otherwise Has clothing section.',
'4. Unreachable (403/429/timeouts) verified rows after 2 retries: kept only if directory/reference category is fashion/department store (signal = directory category).',
'5. Unverified candidates promoted only when live page shows clothing nav terms (>=4 incl. title) + cart evidence + .ae host or AED+971 signals; status marks them as not in the verified directory.',
'6. New discovery: unverified-candidate promotion and a curated list of known UAE-operating fashion brands/retail groups (Inditex, Landmark, Apparel Group, Alshaya etc.) absent from the directory, accepted as Known major retailer – manually accepted (UAE storefronts bot-walled / not live-checkable here).',
'6b. Recall pass: retry of unreachable/blocked stores (30s timeout, low concurrency); Playwright rendering for JS menus (913 thin/excluded directory rows, 329 thin unverified, 581 blocked/thin retries, 200 Common Crawl hosts); Common Crawl cluster.idx/CDX read for *.ae across 3 recent crawls (42k hosts, 270 new fashion-path/name hosts probed); ~315 global brand domains probed for UAE-locale paths (92 hits classified); UAE App Store fashion queries (56 sellerUrls with UAE signals); manual re-check of requested brands. Trendyol removed as a multi-seller marketplace.',
'7. Size tier: Brand-Major = reference-universe major or well-known brand; Large/Medium/Small from Shopify/Woo catalogue size (best effort).',
'Limitations: nav-based keyword signals; JS-rendered sites have thin nav; Gender/Type are heuristic; marketplaces (noon, Amazon, Farfetch, Trendyol-type aggregators) excluded.']: w3.append([t])
w3.column_dimensions['A'].width=120
wb.save('UAE_Fashion_Stores.xlsx')
with open('UAE_Fashion_Stores.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f);w.writerow(H);w.writerows(rows)
print(len(rows))

import pickle,collections,csv,sqlite3,re,openpyxl,urllib.parse as up
from openpyxl.styles import Font,PatternFill
out,excl=pickle.load(open('result.pkl','rb'))
S=pickle.load(open('qa_sample.pkl','rb'))
QA_ERR=[24,67,78]; QA_N=80
tier_ord={'Brand-Major':0,'Large':1,'Medium':2,'Small':3}
out.sort(key=lambda f:(tier_ord[f['tier']],f['name'].lower()))
m=sqlite3.connect('/tmp/sa/data/saudi_ecom.sqlite')
rex=[]
for n,cat,st,u,reason in m.execute("select name,category,status,best_url,reason from reference_stores where (category like '%ashion%' or category like 'Department%' or category like '%Sports%' or category like '%Baby%') and status in ('excluded','screened_out','unresolved')"):
    if any(n==f['name'] for f in out): continue
    rex.append((n,u or '',(reason or '')[:150],cat,'reference'))
head=['Store Name','Website','Type','Gender','Platform','Size Tier','Verification Status','Evidence URL']
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Fashion Stores'; ws.append(head)
for f in out: ws.append([f['name'],f['site'],f['type'],f['gender'],f['platform'],f['tier'],f['status'],f['ev']])
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
for col,w in zip('ABCDEFGH',(36,42,20,20,16,14,36,60)): ws.column_dimensions[col].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
we=wb.create_sheet('Excluded'); we.append(['Store Name','Website','Reason','Stored Category','Key / Source'])
for e in excl: we.append([e[0],e[1],e[2],e[3],e[4]])
for e in rex: we.append(list(e))
for c in we[1]: c.font=Font(bold=True)
cm=wb.create_sheet('Coverage & Method')
def sec(title,items):
    cm.append([title]); cm[cm.max_row][0].font=Font(bold=True)
    for k,v in items: cm.append([k,v])
    cm.append([])
n=len(out)
sec('Totals',[('Fashion stores in final list',n),('Excluded rows listed',len(excl)+len(rex))])
sec('By source',sorted(collections.Counter(f['src'] for f in out).items()))
sec('By type',sorted(collections.Counter(f['type'] for f in out).items()))
g=collections.Counter()
for f in out:
    for x in f['gender'].split(', '): g['covers '+x]+=1
sec('By gender (stores can cover several)',sorted(g.items()))
sec('By gender combination',sorted(collections.Counter(f['gender'] for f in out).items()))
sec('By platform',collections.Counter(f['platform'] for f in out).most_common())
sec('By size tier',sorted(collections.Counter(f['tier'] for f in out).items(),key=lambda x:tier_ord[x[0]]))
sec('By verification status',collections.Counter(f['status'] for f in out).items())
sec('By clothing signal origin',collections.Counter(f['how'] for f in out).items())
sec('Excluded by reason',collections.Counter(re.sub(r'\(.*','',e[2])[:60] for e in excl).most_common(12))
sec('QA',[('Random sample size',QA_N),('Errors found in final audit sample (seed 5)',len(QA_ERR)),('Error rate',f'{len(QA_ERR)/QA_N:.1%}'),('Earlier sample iterations','seed 7: ~10/80 (12.5%) -> rules fixed; seed 99: ~5/80; seed 2026: 6/80 (7.5%) -> kids/equestrian/uniform/embroidery rules fixed; final seed 5: 3/80'),('QA judged by','manual review of store name, category, matched signal (no live re-fetch)')])
sec('Method',[
('1. Pre-classification','All 6,876 verified master-directory rows (incl. secondary-evidence and manual majors) plus ~4.1k unverified candidates (weak Saudi/weak ecommerce/JS/thin) were live-scanned; stored category used only as prior.'),
('2. Live signals','Homepage anchors/links/title/meta (Arabic+English keywords); Shopify /collections.json + /products.json; WooCommerce Store API categories; sitemap.xml category/collection sub-sitemaps; Salla public storefront API (/store/v1/products/categories and /products) used by the storefront JS. Per-host throttled, no CAPTCHA/login bypass, no orders.'),
('3. Acceptance rules','Fashion-category stores need >=1 clothing keyword; other categories need >=2 (>=3 for household/electronics/food-type categories, >=4 for heavy mismatch) distinct clothing terms incl. a clear one; store-name+category hint allowed for fashion categories; kids-only (kids terms without adult/gender terms or kids in store name), medical/uniform/equestrian/embroidery/laundry names excluded.'),
('4. Gender','From women/men/unisex keyword signals in nav/categories + stored category (abaya -> Women, thobe/shemagh -> Men). If no gender word is found the store is labelled Unisex (default, not explicit; count shown above under combinations).'),
('5. Type','Fashion-only = stored fashion category (clothing-led store); Has clothing section = store whose primary category is something else (sports, general, bags, gifts...). Best-effort.'),
('6. Size tier','Brand-Major = reference-universe major or manually accepted brand; Large = reference-universe medium store; Medium = independent custom domain; Small = Salla/Zid/myshopify merchant subdomain. Best-effort.'),
('7. New discovery','(a) 205-URL fashion seed list of GCC/global fashion brands & retailers verified with the existing verifier (119 new candidates, 4 verified, none classified as fashion; ~86 already known); (b) 13 known major fashion brands whose Saudi storefronts are bot-walled/JS-rendered accepted as Known major retailer – manually accepted (reference audit status unresolved); (c) 19 existing unverified candidates with weak Saudi signals only that scanned as fashion (Unverified – weak Saudi signals only). Common Crawl domain-wide keyword filters time out (504); Wayback/CC keyword discovery therefore reused the existing full salla.sa/zid/myshopify/.sa harvest. Salla/Zid showcase and Saudi App Store sellerUrl discovery were already exhausted by the base pipeline.'),
('Limitations','Stores whose storefront could not be fetched (bot walls, timeouts, JS-only Zid pages with no static nav) and no name hint are listed in Excluded as unclassifiable, so recall for small Zid stores is understated. Fetching was from a US VM.')])
cm.column_dimensions['A'].width=48; cm.column_dimensions['B'].width=120
wb.save('/tmp/sa/markets/fashion/Saudi_Arabia_Fashion_Stores.xlsx')
with open('/tmp/sa/markets/fashion/Saudi_Arabia_Fashion_Stores.csv','w',newline='',encoding='utf-8-sig') as fh:
    w=csv.writer(fh); w.writerow(head)
    for f in out: w.writerow([f['name'],f['site'],f['type'],f['gender'],f['platform'],f['tier'],f['status'],f['ev']])
print(n,collections.Counter(f['tier'] for f in out),collections.Counter(f['platform'] for f in out).most_common(6))
print([f['name'] for f in out if f['tier']=='Brand-Major'])

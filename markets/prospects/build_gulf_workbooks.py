"""Build per-market Ecommerce and Fashion prospect workbooks (UAE, Qatar, Kuwait, Saudi Arabia)."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path("/tmp/prospects")
FINAL_DIRS = {"UAE": "final_gulf", "Qatar": "final_gulf", "Kuwait": "final_gulf", "Saudi": "final_saudi"}
FILE_PREFIX = {"Saudi": "Saudi_Arabia"}
OUT = Path(__file__).resolve().parent
NF = "Not found"

DARK = PatternFill("solid", fgColor="1F2933")
ORANGE = "F97316"
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HEADER_BORDER = Border(bottom=Side(style="thick", color=ORANGE))
BODY_FONT = Font(name="Calibri", size=10)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
BAND = PatternFill("solid", fgColor="FFF7ED")

MASTER_COLS = [
    ("Website Name", 34), ("Root Domain Link", 38), ("Semrush Organic Traffic", 16),
    ("Website Count", 13), ("Count Method", 34),
    ("Company Phone Numbers", 36),
    ("Company Emails", 38),
    ("Company LinkedIn", 42), ("Primary Decision Maker Name", 28), ("Position", 40),
]
DM_COLS = [
    ("Company Name", 30), ("Website Link", 38), ("Contact Name", 26), ("Position", 44),
    ("Phone Numbers", 24), ("Email", 30), ("LinkedIn Profile", 50),
]


def clean(v):
    v = "" if v is None else str(v).strip()
    return v if v else NF


def as_https(url):
    return "https://" + url[len("http://"):] if url.startswith("http://") else url


def is_url(v):
    return v.startswith(("http://", "https://"))


def style_sheet(ws, widths, link_cols, num_cols=()):
    for i, (name, w) in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.border = DARK, HEADER_FONT, HEADER_BORDER
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 34
    for r in range(2, ws.max_row + 1):
        for ci in range(1, len(widths) + 1):
            c = ws.cell(row=r, column=ci)
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="right" if ci in num_cols else "left")
            if r % 2 == 1:
                c.fill = BAND
            if ci in link_cols and isinstance(c.value, str) and is_url(c.value):
                c.hyperlink = c.value
                c.font = LINK_FONT
            if ci in num_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{ws.max_row}"
    ws.sheet_properties.tabColor = ORANGE


def build_market(market):
    final = SRC / FINAL_DIRS[market]
    p = pd.read_csv(SRC / f"{market}.csv", dtype=str, keep_default_na=False)
    comp = pd.read_csv(final / f"{market}_companies.csv", dtype=str, keep_default_na=False)
    ppl = pd.read_csv(final / f"{market}_people.csv", dtype=str, keep_default_na=False)
    assert p["Root Domain"].is_unique and comp["Root Domain"].is_unique
    assert set(p["Root Domain"]) == set(comp["Root Domain"])
    df = p.merge(comp, on="Root Domain", how="left", validate="one_to_one")
    df["traffic"] = pd.to_numeric(df["Semrush Organic Traffic"], errors="coerce").fillna(0)
    df = df.sort_values("traffic", ascending=False, kind="stable")
    results = []
    for kind, tag in (("Ecommerce", "directory"), ("Fashion", "fashion")):
        sel = df[df["Lists"].str.split(";").apply(lambda x: tag in x)]
        domains = set(sel["Root Domain"])
        wb = Workbook()
        ws = wb.active
        ws.title = "Master Prospects"
        ws.append([c for c, _ in MASTER_COLS])
        for _, r in sel.iterrows():
            count = r["Website Count"].strip()
            ws.append([
                clean(r["Website Name"]), as_https(r["Root Domain"]), int(r["traffic"]),
                int(count) if count.isdigit() else NF, clean(r["Count Method"]),
                clean(r["Company Phone Numbers"]), clean(r["Company Emails"]),
                clean(r["Company LinkedIn"]), clean(r["Primary Decision Maker Name"]),
                clean(r["Position_y"] if "Position_y" in r else r["Position"]),
            ])
        style_sheet(ws, MASTER_COLS, link_cols={2, 8}, num_cols={3, 4})
        pp = ppl[ppl["Website Link"].isin(domains)]
        order = {d: i for i, d in enumerate(sel["Root Domain"])}
        pp = pp.assign(_o=pp["Website Link"].map(order)).sort_values("_o", kind="stable")
        ws2 = wb.create_sheet("Decision Makers")
        ws2.append([c for c, _ in DM_COLS])
        for _, r in pp.iterrows():
            ws2.append([
                clean(r["Company Name"]), as_https(r["Website Link"]), clean(r["Contact Name"]),
                clean(r["Position"]), clean(r["Phone Numbers"]), clean(r["Email"]),
                clean(r["LinkedIn Profile"]),
            ])
        style_sheet(ws2, DM_COLS, link_cols={2, 7})
        path = OUT / f"{FILE_PREFIX.get(market, market)}_{kind}_Prospects.xlsx"
        wb.save(path)
        results.append((path, len(sel), len(pp), set(sel["Root Domain"])))
    return results


def validate(path, n_master, n_dm, domains_http):
    wb = load_workbook(path)
    assert wb.sheetnames == ["Master Prospects", "Decision Makers"], wb.sheetnames
    m, d = wb["Master Prospects"], wb["Decision Makers"]
    assert m.max_row - 1 == n_master and d.max_row - 1 == n_dm
    for ws in (m, d):
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref
        for row in ws.iter_rows(min_row=1):
            for c in row:
                assert c.value is not None and str(c.value).strip() != "", (path.name, c.coordinate)
    links = [m.cell(r, 2).value for r in range(2, m.max_row + 1)]
    assert len(set(links)) == len(links) == n_master
    assert all(l.startswith("https://") and m.cell(i + 2, 2).hyperlink for i, l in enumerate(links))
    traffic = [m.cell(r, 3).value for r in range(2, m.max_row + 1)]
    assert traffic == sorted(traffic, reverse=True)
    dm_sites = {d.cell(r, 2).value for r in range(2, d.max_row + 1)}
    assert dm_sites <= set(links)


if __name__ == "__main__":
    for market in sys.argv[1:] or ["UAE", "Qatar", "Kuwait", "Saudi"]:
        for path, nm, nd, doms in build_market(market):
            validate(path, nm, nd, doms)
            print(f"{path.name}: Master Prospects={nm} rows, Decision Makers={nd} rows  [OK]")
